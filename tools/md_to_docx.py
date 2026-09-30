#!/usr/bin/env python3
"""Convert a book-style Markdown file into a styled Word (.docx) document.

Usage: python tools/md_to_docx.py docs/PROJECT_BOOK.md docs/AI-Productivity-Hub-Guide.docx
"""
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

TOKEN = re.compile(r'(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))')

HEADING = {
    1: 'Heading 1',
    2: 'Heading 2',
    3: 'Heading 3',
    4: 'Heading 4',
}

ACCENT = RGBColor(0x1F, 0x49, 0x7D)
CODE_BG = 'F2F2F2'


def shade(paragraph, color=CODE_BG):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), color)
    pPr.append(shd)


def add_runs(paragraph, text):
    for chunk in TOKEN.split(text):
        if not chunk:
            continue
        if chunk.startswith('**') and chunk.endswith('**'):
            run = paragraph.add_run(chunk[2:-2])
            run.bold = True
        elif chunk.startswith('`') and chunk.endswith('`'):
            run = paragraph.add_run(chunk[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
        elif re.fullmatch(r'\[[^\]]+\]\([^)]+\)', chunk):
            paragraph.add_run(re.match(r'\[([^\]]+)\]', chunk).group(1))
        else:
            paragraph.add_run(chunk)


def table_widths(table, total=6.5):
    table.autofit = True
    cols = max(len(row.cells) for row in table.rows)
    width = Inches(total / max(cols, 1))
    for row in table.rows:
        for cell in row.cells:
            cell.width = width


def add_table(doc, rows):
    if not rows:
        return
    header = [c.strip() for c in rows[0].split('|')[1:-1]]
    body = [[c.strip() for c in row.split('|')[1:-1]] for row in rows[2:]]

    table = doc.add_table(rows=1 + len(body), cols=len(header))
    table.style = 'Light Grid Accent 1'
    table_widths(table)

    for j, text in enumerate(header):
        cell = table.rows[0].cells[j]
        cell.paragraphs[0].text = ''
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for i, row in enumerate(body, start=1):
        for j, text in enumerate(row):
            p = table.rows[i].cells[j].paragraphs[0]
            add_runs(p, text)


def main(src, dst):
    doc = Document()

    title = doc.add_heading('AI-Productivity-Hub — The Complete Project Guide', 0)
    for run in title.runs:
        run.font.color.rgb = ACCENT

    doc.add_paragraph(
        'A plain-English, book-style walkthrough of everything in this project: '
        'what it does, how it is built, how the pieces fit together, and how to '
        'run, secure and operate it.'
    )

    with open(src, encoding='utf-8') as f:
        lines = f.read().splitlines()

    i = 0
    in_code = False
    code_lines = []
    table_rows = []

    def flush_code():
        nonlocal code_lines
        if not code_lines:
            return
        for cl in code_lines:
            p = doc.add_paragraph()
            run = p.add_run(cl)
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
            p.paragraph_format.left_indent = Pt(18)
            p.paragraph_format.space_after = Pt(0)
            shade(p)
        doc.add_paragraph()

    def flush_table():
        nonlocal table_rows
        if table_rows:
            add_table(doc, table_rows)
            doc.add_paragraph()
            table_rows = []

    while i < len(lines):
        line = lines[i]

        if line.startswith('```'):
            if in_code:
                flush_code()
                code_lines = []
                in_code = False
            else:
                flush_table()
                in_code = True
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if line.strip() == '':
            flush_table()
            i += 1
            continue

        if line.startswith('|'):
            table_rows.append(line)
            i += 1
            continue

        flush_table()

        stripped = line.strip()

        if re.fullmatch(r'-{3,}', stripped):
            i += 1
            continue

        m = re.match(r'^(#{1,4})\s+(.*)$', stripped)
        if m:
            level = len(m.group(1))
            p = doc.add_paragraph()
            p.style = doc.styles[HEADING[level]]
            add_runs(p, m.group(2))
            i += 1
            continue

        if stripped.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            add_runs(p, stripped[2:])
            i += 1
            continue

        m = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if m:
            p = doc.add_paragraph(style='List Number')
            add_runs(p, m.group(2))
            i += 1
            continue

        p = doc.add_paragraph()
        add_runs(p, stripped)
        i += 1

    flush_code()
    flush_table()

    doc.save(dst)
    print(f"Saved: {dst}")


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])